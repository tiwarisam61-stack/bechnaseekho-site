import re
from pathlib import Path
from datetime import date

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "requirements" / "CareerSync_Stakeholder_Website_Walkthrough.pptx"

INDEX = ROOT / "index.html"
SERVER = ROOT / "server.js"
ADMIN_DASH = ROOT / "admin-dashboard.html"
JOB_DETAILS = ROOT / "job-details.html"
MY_APPS = ROOT / "my-applications.html"
ADMIN_APPS = ROOT / "admin-applications.html"
README = ROOT / "README.md"


def read_text(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="ignore")


def clean_html_text(raw: str) -> str:
    t = re.sub(r"<[^>]+>", " ", raw)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def add_title_slide(prs: Presentation, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs: Presentation, title: str, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    for i, b in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = b
        p.level = 0
        p.font.size = Pt(18)


def add_two_col_slide(prs: Presentation, title: str, left_title: str, left_lines, right_title: str, right_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.6))
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.4), Inches(6.1), Inches(5.6))

    left_tf = left_box.text_frame
    right_tf = right_box.text_frame

    left_tf.word_wrap = True
    right_tf.word_wrap = True

    lp = left_tf.paragraphs[0]
    lp.text = left_title
    lp.font.bold = True
    lp.font.size = Pt(22)

    for line in left_lines:
        p = left_tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(16)

    rp = right_tf.paragraphs[0]
    rp.text = right_title
    rp.font.bold = True
    rp.font.size = Pt(22)

    for line in right_lines:
        p = right_tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(16)


def parse_routes(server_text: str):
    return re.findall(r"app\.(get|post)\('([^']+)'", server_text)


def parse_action_labels(index_text: str):
    m = re.search(r"const m = \{(.+?)\};", index_text, re.S)
    if not m:
        return []
    body = m.group(1)
    keys = re.findall(r'"([^"]+)"\s*:\s*\[', body)
    return sorted(set(keys))


def parse_coming_soon_labels(index_text: str):
    m = re.search(r"comingSoonLabels\s*=\s*new Set\(\[(.*?)\]\)", index_text, re.S)
    if not m:
        return []
    values = re.findall(r"'([^']+)'", m.group(1))
    return values


def parse_nav_labels(index_text: str):
    labels = re.findall(r'<a class="nav-link[^>]*>(.*?)</a>', index_text, re.S)
    clean = [clean_html_text(x) for x in labels]
    return [x for x in clean if x]


def parse_resource_titles(index_text: str):
    return re.findall(r'<div class="resource-title">(.*?)</div>', index_text, re.S)


def parse_admin_sections(admin_text: str):
    items = re.findall(r"SECTION: ([A-Z\s]+)", admin_text)
    return [x.title() for x in items]


def first_n(lines, n=10):
    return lines[:n] if len(lines) > n else lines


def main():
    index_text = read_text(INDEX)
    server_text = read_text(SERVER)
    admin_text = read_text(ADMIN_DASH)
    job_text = read_text(JOB_DETAILS)
    my_text = read_text(MY_APPS)
    admin_apps_text = read_text(ADMIN_APPS)
    readme_text = read_text(README)

    routes = parse_routes(server_text)
    action_labels = parse_action_labels(index_text)
    coming_soon = parse_coming_soon_labels(index_text)
    nav_labels = parse_nav_labels(index_text)
    resource_titles = [clean_html_text(x) for x in parse_resource_titles(index_text)]
    admin_sections = parse_admin_sections(admin_text)

    total_buttons = len(re.findall(r"<button\b", index_text, re.I))
    total_onclick = len(re.findall(r"onclick=", index_text, re.I))

    prs = Presentation()
    prs.core_properties.title = "CareerSync Website Walkthrough"
    prs.core_properties.subject = "UI map, button map, feature map, and architecture overview"
    prs.core_properties.author = "CareerSync Dev Team"

    add_title_slide(
        prs,
        "CareerSync Website - Stakeholder Walkthrough",
        f"Generated on {date.today().isoformat()} | UI, Buttons, Flows, APIs, Admin",
    )

    add_bullet_slide(
        prs,
        "1. Project Snapshot",
        [
            "Application: CareerSync Web App",
            "Main user pages: index, job-details, my-applications",
            "Admin pages: admin-dashboard, admin-applications",
            "Backend: Node.js + Express with JSON data and file uploads",
            "Goal of this deck: explain what exists, where it lives, and what each click does",
        ],
    )

    add_bullet_slide(
        prs,
        "2. Tech Stack",
        [
            "Frontend: HTML, CSS, JavaScript in page files",
            "Backend: Node.js + Express (server.js)",
            "Security + infra: helmet, compression, express-rate-limit",
            "File upload: multer (resume upload endpoint)",
            "Data: data/jobs.json and data/applications.json",
        ],
    )

    add_bullet_slide(
        prs,
        "3. Folder Structure",
        [
            "assets/company-logos/ - company logo assets",
            "assets/images/branding/ - brand logo",
            "assets/images/tools/ - menu tool icons",
            "assets/images/resources/ - resources section images",
            "data/ - jobs + applications JSON",
            "uploads/ - uploaded resumes",
            "docs/requirements/ - project documentation",
        ],
    )

    add_bullet_slide(
        prs,
        "4. Main Navigation and User Entry Points",
        [
            "Top navigation labels: " + ", ".join(nav_labels),
            "Career Tools and Resources open dropdown panels",
            "Profile menu includes My Profile, Dashboard, My Applications, Settings, Admin Panel, Logout",
            "Mobile drawer mirrors desktop navigation and dropdown options",
            "Newsletter form and footer links provide additional entry points",
        ],
    )

    add_two_col_slide(
        prs,
        "5. Action Handler Map (index.html)",
        "Clickable Labels Routed by handleAction(label)",
        [f"- {x}" for x in first_n(action_labels, 12)] + ([f"- ... and {max(len(action_labels)-12, 0)} more"] if len(action_labels) > 12 else []),
        "Current Coming Soon Labels",
        [f"- {x}" for x in coming_soon],
    )

    add_bullet_slide(
        prs,
        "6. Career Tools Dropdown (User Experience)",
        [
            "AI Job Matcher - Coming Soon",
            "Resume Builder - visible in menu, handled as Coming Soon",
            "ATS Score Checker - Coming Soon",
            "Mock Interview - Coming Soon",
            "AI Career Coach - Coming Soon",
            "Salary Guide - Coming Soon",
        ],
    )

    add_bullet_slide(
        prs,
        "7. Resources Dropdown and Resources Section",
        [
            "Dropdown options: Career Blog, Interview Questions, Resume Templates, Salary Calculator, Career Tips",
            "Resources grid cards are present with image + title + short description",
            "Resource cards currently show Coming Soon badges",
            "Example resource titles:",
        ] + [f"- {x}" for x in first_n(resource_titles, 6)],
    )

    add_bullet_slide(
        prs,
        "8. Job Details Page Flow",
        [
            "Route entry: /jobs/:jobId or job-details.html",
            "Loads job metadata from /api/jobs/:jobId",
            "Application form validates required fields and file constraints",
            "Duplicate-check endpoint used before submit",
            "Submit posts multipart form to /api/applications",
            "Sticky Apply button and WhatsApp support action are available",
        ],
    )

    add_bullet_slide(
        prs,
        "9. My Applications Page Flow",
        [
            "Page title: My Applications",
            "Fetches user applications via /api/applications?email=...",
            "Displays status/history per application",
            "Back-to-home navigation to main landing page",
        ],
    )

    add_bullet_slide(
        prs,
        "10. Admin Applications Page",
        [
            "Page title: Admin Applications",
            "Capabilities: filter/search, status update, delete record",
            "Export actions: CSV and JSON",
            "APIs used: /api/applications and /api/applications/export.*",
        ],
    )

    add_bullet_slide(
        prs,
        "11. Admin Dashboard Modules",
        [
            "Login with role-based access (Super Admin / Manager)",
            "Sidebar module navigation across operational sections",
            "Top-level modules include Home, Applications, Resumes, Candidates, Analytics, Jobs, Companies",
            "Additional modules include Managers, Notifications, Reports, Settings",
            "Charts and KPI widgets are rendered in dashboard sections",
        ] + [f"- Section in code: {x}" for x in first_n(admin_sections, 10)],
    )

    route_lines = [f"{method.upper()} {route}" for method, route in routes]
    add_bullet_slide(
        prs,
        "12. Backend Route Inventory",
        first_n(route_lines, 14) + ([f"... and {max(len(route_lines)-14, 0)} more routes"] if len(route_lines) > 14 else []),
    )

    add_bullet_slide(
        prs,
        "13. API and Data Responsibilities",
        [
            "GET /api/jobs - list available jobs",
            "GET /api/jobs/:jobId - single job details",
            "POST /api/applications - create application + resume upload",
            "GET /api/applications - list applications",
            "GET /api/applications/export.csv and export.json - admin export",
            "Data files: data/jobs.json and data/applications.json",
        ],
    )

    add_bullet_slide(
        prs,
        "14. Security and Stability Controls",
        [
            "Helmet headers enabled",
            "Compression enabled",
            "Rate limiting on /api and /api/applications",
            "Upload file type and size restrictions for resumes",
            "Ready and health endpoints available for monitoring",
        ],
    )

    add_two_col_slide(
        prs,
        "15. UI Interaction Metrics (Current Code)",
        "index.html Interaction Footprint",
        [
            f"- Total button tags: {total_buttons}",
            f"- Total onclick attributes: {total_onclick}",
            "- Central interaction router: handleAction(label)",
            "- Nav scrolling handler: handleNavClick(event, target)",
            "- Coming-soon protection click interceptor enabled",
        ],
        "Main User Journeys",
        [
            "- Browse jobs from homepage",
            "- Open job details and submit application",
            "- View application history in My Applications",
            "- Admin login and manage applications",
            "- Use dropdowns for Career Tools and Resources",
        ],
    )

    add_bullet_slide(
        prs,
        "16. Button-to-Outcome Summary (Stakeholder View)",
        [
            "Navbar buttons -> smooth scroll to sections",
            "Dropdown tool/resource buttons -> currently Coming Soon where marked",
            "Apply Now buttons -> open job details/application workflow",
            "My Applications -> reads status from backend",
            "Admin buttons -> section switching and CRUD/export operations",
            "Newsletter submit -> communication workflow handlers",
        ],
    )

    add_bullet_slide(
        prs,
        "17. Code Ownership Map",
        [
            "UI + interactions: index.html",
            "Apply workflow: job-details.html + /api/applications",
            "Candidate tracking UI: my-applications.html",
            "Admin management UI: admin-dashboard.html + admin-applications.html",
            "Server and API routes: server.js",
            "Config/deps: package.json",
        ],
    )

    add_bullet_slide(
        prs,
        "18. Readiness for New Development",
        [
            "Core pages are structured and separated by function",
            "Assets are organized under assets/images and assets/company-logos",
            "Documentation exists in README and docs/requirements",
            "APIs and routes are in one backend entry file",
            "Recommended next step: feature roadmap and test-case deck extension",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"PPT generated: {OUT}")


if __name__ == "__main__":
    main()
