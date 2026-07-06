import re
from pathlib import Path
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "requirements" / "CareerSync_Stakeholder_Deck_Beautiful.pptx"

INDEX = ROOT / "index.html"
SERVER = ROOT / "server.js"
ADMIN_DASH = ROOT / "admin-dashboard.html"
JOB_DETAILS = ROOT / "job-details.html"
README = ROOT / "README.md"

SS = ROOT / "docs" / "requirements" / "screenshots"
LOGO = ROOT / "assets" / "images" / "branding" / "nav-logo.png"


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def add_bg(slide, rgb=(245, 247, 255)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def add_header(slide, title: str, subtitle: str | None = None):
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9.8), Inches(0.7))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = RGBColor(34, 40, 58)
    if subtitle:
        sp = tf.add_paragraph()
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.color.rgb = RGBColor(94, 106, 132)

    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(11.5), Inches(0.15), width=Inches(1.1))


def add_bullets(slide, x, y, w, h, lines, font_size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(40, 45, 62)


def add_image(slide, img_path: Path, x, y, w=None, h=None):
    if not img_path.exists():
        return
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    slide.shapes.add_picture(str(img_path), Inches(x), Inches(y), **kw)


def clean_html_text(raw: str) -> str:
    t = re.sub(r"<[^>]+>", " ", raw)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def parse_routes(server_text: str):
    return re.findall(r"app\.(get|post)\('([^']+)'", server_text)


def parse_action_labels(index_text: str):
    m = re.search(r"const m = \{(.+?)\};", index_text, re.S)
    if not m:
        return []
    body = m.group(1)
    return sorted(set(re.findall(r'"([^"]+)"\s*:\s*\[', body)))


def parse_nav(index_text: str):
    labels = re.findall(r'<a class="nav-link[^>]*>(.*?)</a>', index_text, re.S)
    return [clean_html_text(x) for x in labels if clean_html_text(x)]


def parse_coming_soon(index_text: str):
    m = re.search(r"comingSoonLabels\s*=\s*new Set\(\[(.*?)\]\)", index_text, re.S)
    if not m:
        return []
    return re.findall(r"'([^']+)'", m.group(1))


def parse_admin_sections(admin_text: str):
    return [x.title() for x in re.findall(r"SECTION: ([A-Z\s]+)", admin_text)]


def cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, (233, 238, 255))
    add_header(slide, "CareerSync Website Walkthrough", f"Stakeholder Deck | {date.today().isoformat()}")
    add_bullets(
        slide,
        0.7,
        1.4,
        5.6,
        2.0,
        [
            "UI + Button map + Flow map + Technical architecture",
            "Prepared from current live codebase and runtime structure",
        ],
        font_size=17,
    )
    add_image(slide, SS / "index-home-section.png", 6.0, 1.2, w=6.8)


def content_slide(prs, title, subtitle, bullets=None, image: Path | None = None, image_pos=(7.0, 1.3, 5.8, 4.9)):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, title, subtitle)
    if bullets:
        add_bullets(slide, 0.7, 1.4, 6.0, 5.7, bullets)
    if image:
        x, y, w, h = image_pos
        add_image(slide, image, x, y, w=w, h=h)


def two_image_slide(prs, title, subtitle, left_img: Path, right_img: Path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, title, subtitle)
    add_image(slide, left_img, 0.7, 1.4, w=6.0, h=5.4)
    add_image(slide, right_img, 6.9, 1.4, w=6.0, h=5.4)


def button_gallery_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header(slide, "Button Gallery", "Actual captured buttons used in user/admin workflows")

    imgs = [
        SS / "job-details-apply-button.png",
        SS / "job-details-submit-button.png",
        SS / "admin-export-csv-button.png",
        SS / "admin-export-json-button.png",
    ]
    positions = [
        (0.8, 1.5, 3.0, 1.0),
        (4.0, 1.5, 3.2, 1.0),
        (7.4, 1.5, 2.6, 1.0),
        (10.2, 1.5, 2.6, 1.0),
    ]
    for img, pos in zip(imgs, positions):
        x, y, w, h = pos
        add_image(slide, img, x, y, w=w, h=h)

    add_bullets(
        slide,
        0.8,
        3.0,
        12.0,
        3.6,
        [
            "Apply Now: opens/sticks job application flow.",
            "Submit Application: posts validated form data to backend.",
            "Export CSV / JSON: admin download actions for reporting.",
            "All critical buttons are wired to explicit handlers or API calls.",
        ],
        font_size=16,
    )


def build():
    index_text = read_text(INDEX)
    server_text = read_text(SERVER)
    admin_text = read_text(ADMIN_DASH)
    _ = read_text(JOB_DETAILS)
    _ = read_text(README)

    routes = [f"{m.upper()} {r}" for m, r in parse_routes(server_text)]
    actions = parse_action_labels(index_text)
    nav = parse_nav(index_text)
    soon = parse_coming_soon(index_text)
    admin_sections = parse_admin_sections(admin_text)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    cover_slide(prs)

    content_slide(
        prs,
        "Executive Summary",
        "What this website currently provides",
        [
            "CareerSync is a full web platform for jobs, applications, and admin management.",
            "Frontend pages are separated by role: user journey and admin journey.",
            "Backend serves APIs, static assets, file uploads, exports, and health endpoints.",
            "UI has structured sections, dropdown tools/resources, and form validation.",
        ],
        SS / "index-navbar.png",
    )

    content_slide(
        prs,
        "Technology Stack",
        "Languages, framework, and runtime",
        [
            "Frontend language: HTML + CSS + JavaScript.",
            "Backend language: JavaScript (Node.js + Express).",
            "Data layer: JSON files for jobs and applications.",
            "Upload handling: Multer with size/type checks.",
            "Security/middleware: Helmet, rate-limit, compression.",
        ],
        SS / "index-home.png",
        image_pos=(6.8, 1.2, 6.0, 5.8),
    )

    content_slide(
        prs,
        "Website Structure",
        "Main folders and file responsibilities",
        [
            "index.html -> main UI and client interaction logic",
            "job-details.html -> details page + apply form",
            "my-applications.html -> user application history",
            "admin-dashboard.html -> admin control center",
            "admin-applications.html -> application operations + export",
            "server.js -> all routes and backend workflow",
        ],
        SS / "index-footer.png",
    )

    content_slide(
        prs,
        "Main Navigation Map",
        "Top menu, profile, and mobile flows",
        [
            "Top labels: " + ", ".join(nav),
            "Career Tools and Resources use dropdown-based navigation.",
            "Profile menu opens My Profile, Dashboard, My Applications, Settings, Admin Panel.",
            "Mobile drawer mirrors desktop navigation and action options.",
        ],
        SS / "index-navbar.png",
    )

    content_slide(
        prs,
        "Action Router and Click Behavior",
        "Central mapping handled in index.html",
        [
            "Main click router: handleAction(label).",
            "Mapped labels include login, signup, notifications, contact, and utility actions.",
            "Resume/resource buttons currently handled as Coming Soon where configured.",
            "Detected mapped action labels (sample):",
        ] + [f"- {x}" for x in actions[:10]],
        SS / "index-services-bento.png",
    )

    content_slide(
        prs,
        "Coming Soon Control",
        "Protected options and UX behavior",
        [
            "Configured Coming Soon labels:",
        ] + [f"- {x}" for x in soon],
        SS / "index-resources-grid.png",
    )

    content_slide(
        prs,
        "Career Tools Experience",
        "Visual section for AI tools and cards",
        [
            "Bento layout highlights tools and primary actions.",
            "Some options are active, many are safely marked Coming Soon.",
            "Badges and disabled interactions avoid broken user journeys.",
            "Footer CTA in tools panel is intentionally blocked for upcoming releases.",
        ],
        SS / "index-services-bento.png",
    )

    content_slide(
        prs,
        "Resources Experience",
        "Templates, interview prep, blog, salary guide, and tips",
        [
            "Resource cards include image, tag, title, description, and CTA text.",
            "Current resource options are visually complete and structured.",
            "Resume Templates card is visible as requested.",
        ],
        SS / "index-resource-card-template.png",
        image_pos=(7.2, 1.4, 5.5, 5.3),
    )

    content_slide(
        prs,
        "Job Details Page",
        "From job listing to complete application submission",
        [
            "Displays salary, location, role data, and recruiter/company context.",
            "Apply form validates required fields, email/phone, resume file size/type.",
            "Uses duplicate-check endpoint before final submit.",
            "Sticky Apply + WhatsApp support buttons improve conversion.",
        ],
        SS / "job-details-page.png",
    )

    button_gallery_slide(prs)

    content_slide(
        prs,
        "My Applications",
        "Candidate tracking view",
        [
            "Fetches records for user email via API.",
            "Shows application history and statuses.",
            "Simple and focused user follow-up page.",
        ],
        SS / "my-applications-page.png",
    )

    two_image_slide(
        prs,
        "Admin Entry and Control Center",
        "Login experience and main dashboard",
        SS / "admin-dashboard-login.png",
        SS / "admin-dashboard-main.png",
    )

    content_slide(
        prs,
        "Admin Applications Operations",
        "Review, status update, delete, and export",
        [
            "Search and filter over application records.",
            "Update status workflow from admin table.",
            "Delete action for data cleanup.",
            "One-click CSV and JSON exports for reports.",
        ],
        SS / "admin-applications-page.png",
    )

    content_slide(
        prs,
        "Admin Dashboard Sections",
        "Role-based module navigation",
        [
            "Key sections include:",
        ] + [f"- {s}" for s in admin_sections[:12]],
        SS / "admin-dashboard-main.png",
    )

    content_slide(
        prs,
        "Backend Route Inventory",
        "Express routes currently available",
        [f"- {r}" for r in routes[:16]] + ([f"- ... and {len(routes) - 16} more"] if len(routes) > 16 else []),
        SS / "job-details-page.png",
        image_pos=(7.1, 1.4, 5.7, 5.3),
    )

    content_slide(
        prs,
        "API Flow and Data",
        "How user/admin UI talks to backend",
        [
            "GET /api/jobs and GET /api/jobs/:jobId drive listing/details.",
            "POST /api/applications handles multipart form submit.",
            "GET /api/applications supports user/admin listing views.",
            "Export endpoints provide reporting downloads.",
            "Source files: data/jobs.json and data/applications.json.",
        ],
        SS / "admin-applications-page.png",
    )

    content_slide(
        prs,
        "Branding and Asset Strategy",
        "Logo, tool icons, and resources images",
        [
            "Brand logo centralized in assets/images/branding.",
            "Tool icons centralized in assets/images/tools.",
            "Resource images centralized in assets/images/resources.",
            "Folder strategy supports easy replacement for designers/developers.",
        ],
        LOGO,
        image_pos=(8.5, 2.0, 3.2, 3.2),
    )

    content_slide(
        prs,
        "Security, Stability, and Ops",
        "Production safety controls",
        [
            "Helmet security headers enabled.",
            "Rate limits on API and apply endpoints.",
            "Multer enforces upload limits and file types.",
            "Health and readiness endpoints included.",
            "Static serving with caching enabled.",
        ],
        SS / "index-home.png",
        image_pos=(7.0, 1.3, 5.8, 5.7),
    )

    content_slide(
        prs,
        "Stakeholder FAQ Slide",
        "Quick answers for demo meeting",
        [
            "Q: Which code language is used? A: Frontend HTML/CSS/JS + Backend Node/Express.",
            "Q: Which button opens what? A: Mapped by handleAction and page-specific handlers.",
            "Q: Where are admin controls? A: admin-dashboard and admin-applications pages.",
            "Q: Where are images and logos? A: assets/images and assets/company-logos.",
            "Q: Is data/API wired? A: Yes, jobs/applications APIs and exports are available.",
        ],
        SS / "index-navbar.png",
    )

    content_slide(
        prs,
        "Final Readiness",
        "Website status for stakeholder presentation",
        [
            "UI structure is clear and role-based.",
            "Button behaviors are mapped and documented.",
            "Core user/admin workflows are functional.",
            "Visual evidence included from live screenshots.",
            "Deck is ready to present to stakeholders.",
        ],
        SS / "index-home-section.png",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Generated: {OUT}")


if __name__ == "__main__":
    build()
